"""
月彙總報表 + 自動寄送 Email

讀取所有商品的「日」與「週」三關價分析結果，依「月份」分頁彙整成一份 Excel，
分析完成後自動寄送到指定信箱。

每個月份會產出兩個分頁：
    2026-09      該月每個交易日、每個商品的日線關卡
    2026-09-週   該月每一週、每個商品的週線關卡

執行時機：three_gate/src/analyze.py 之後執行（GitHub Actions 內串接）

寄信需要的環境變數（GitHub Secrets）：
    SMTP_USER  寄件 Gmail 帳號
    SMTP_PASS  Gmail 應用程式密碼 (App Password，不是登入密碼)
    MAIL_TO    收件信箱（不設定則預設寄給 SMTP_USER 自己）
"""

import os
import smtplib
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

import sys
sys.path.insert(0, str(Path(__file__).parent))

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from analyze import load_symbols, analyze_one


ROOT = Path(__file__).parent.parent
REPORT_DIR = ROOT / 'reports'
REPORT_PATH = REPORT_DIR / 'monthly_report.xlsx'

HEADER_BG = '4472C4'
S_BG = 'C6EFCE'
B_BG = 'FFC7CE'
M_BG = 'FFF2CC'

DAILY_HEADERS = [
    '日期', '代碼', '商品', '開', '高', '低', '收', '漲跌', '漲%',
    'S3', 'S2', 'S1', 'M', 'B1', 'B2', 'B3', '方向', '警訊',
]

WEEKLY_HEADERS = [
    '週結算日', '代碼', '商品', '開', '高', '低', '收', '週漲跌', '週漲%',
    'S3', 'S2', 'S1', 'M', 'B1', 'B2', 'B3', '方向', '警訊',
]


def collect_rows(period: str) -> pd.DataFrame:
    """跑一次所有商品的指定週期（daily/weekly）分析，把每列資料收集起來"""
    cfg = load_symbols()
    items = [(it, 'manual') for it in cfg.get('manual', [])] + \
            [(it, 'auto') for it in cfg.get('auto', [])]

    all_rows = []
    for item, source in items:
        code = item['code']
        name = item['name']
        a = analyze_one(code, name, source, period)
        if not a:
            continue
        df = a['df'].copy()
        df['code'] = code
        df['name'] = name
        all_rows.append(df)

    if not all_rows:
        return pd.DataFrame()

    full = pd.concat(all_rows, ignore_index=True)
    full['date'] = pd.to_datetime(full['date'])
    full['year_month'] = full['date'].dt.strftime('%Y-%m')
    return full


def _write_header(ws, headers):
    font_h = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    align_c = Alignment(horizontal='center', vertical='center')

    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c)
        cell.value = h
        cell.font = font_h
        cell.alignment = align_c
        if h in ('S3', 'S2', 'S1'):
            cell.fill = PatternFill('solid', start_color=S_BG)
        elif h in ('B1', 'B2', 'B3'):
            cell.fill = PatternFill('solid', start_color=B_BG)
        elif h == 'M':
            cell.fill = PatternFill('solid', start_color=M_BG)
        else:
            cell.fill = PatternFill('solid', start_color=HEADER_BG)


def _write_rows(ws, sub: pd.DataFrame, headers):
    font_red = Font(name='Calibri', size=10, color='C00000', bold=True)
    font_green = Font(name='Calibri', size=10, color='006100', bold=True)
    align_c = Alignment(horizontal='center', vertical='center')

    sub = sub.sort_values(['date', 'code'])

    r = 2
    for _, row in sub.iterrows():
        ws.cell(r, 1).value = row['date'].strftime('%Y-%m-%d')
        ws.cell(r, 2).value = row['code']
        ws.cell(r, 3).value = row['name']
        for c, key in enumerate(['open', 'high', 'low', 'close', 'change'], 4):
            ws.cell(r, c).value = round(float(row[key]), 2)
        ws.cell(r, 9).value = round(float(row['change_pct']), 4)
        ws.cell(r, 9).number_format = '+0.00%;-0.00%;0.00%'
        for c, key in enumerate(['S3', 'S2', 'S1', 'M', 'B1', 'B2', 'B3'], 10):
            ws.cell(r, c).value = round(float(row[key]), 2)
        dcell = ws.cell(r, 17)
        dcell.value = row['direction']
        dcell.font = font_red if row['direction'] == '多方' else font_green
        dcell.alignment = align_c
        ws.cell(r, 18).value = row['warning'] or ''
        r += 1

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10
    ws.freeze_panes = 'A2'


def write_monthly_workbook(daily_full: pd.DataFrame, weekly_full: pd.DataFrame, out_path: Path):
    wb = Workbook()
    wb.remove(wb.active)  # 移除預設空白頁

    daily_months = set(daily_full['year_month'].unique()) if not daily_full.empty else set()
    weekly_months = set(weekly_full['year_month'].unique()) if not weekly_full.empty else set()
    months = sorted(daily_months | weekly_months)

    for ym in months:
        # 日線分頁（維持原本命名，例如 "2026-09"）
        if ym in daily_months:
            ws = wb.create_sheet(title=ym)
            _write_header(ws, DAILY_HEADERS)
            _write_rows(ws, daily_full[daily_full['year_month'] == ym], DAILY_HEADERS)

        # 週線分頁（例如 "2026-09-週"）
        if ym in weekly_months:
            ws = wb.create_sheet(title=f'{ym}-週')
            _write_header(ws, WEEKLY_HEADERS)
            _write_rows(ws, weekly_full[weekly_full['year_month'] == ym], WEEKLY_HEADERS)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def send_email(attachment_path: Path):
    smtp_user = os.environ.get('SMTP_USER')
    smtp_pass = os.environ.get('SMTP_PASS')
    mail_to = os.environ.get('MAIL_TO') or smtp_user

    if not smtp_user or not smtp_pass:
        print('未設定 SMTP_USER / SMTP_PASS，略過寄信（報表仍會產出並提交到 repo）')
        return

    today = datetime.now().strftime('%Y-%m-%d')
    msg = EmailMessage()
    msg['Subject'] = f'三關價每日彙總報表 {today}'
    msg['From'] = smtp_user
    msg['To'] = mail_to
    msg.set_content(
        f'附件為 {today} 三關價每日彙總報表（依月份分頁，含日線與週線，全商品彙整），系統自動寄送。'
    )

    with open(attachment_path, 'rb') as f:
        data = f.read()
    msg.add_attachment(
        data,
        maintype='application',
        subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        filename=attachment_path.name,
    )

    with smtplib.SMTP_SSL('smtp.gmail.com', 465) as server:
        server.login(smtp_user, smtp_pass)
        server.send_message(msg)

    print(f'已寄出 email 給 {mail_to}')


def main():
    daily_full = collect_rows('daily')
    weekly_full = collect_rows('weekly')

    if daily_full.empty and weekly_full.empty:
        print('沒有任何商品資料，略過月彙總報表')
        return

    write_monthly_workbook(daily_full, weekly_full, REPORT_PATH)
    print(f'✓ 產出月彙總報表: {REPORT_PATH}')

    send_email(REPORT_PATH)


if __name__ == '__main__':
    main()
